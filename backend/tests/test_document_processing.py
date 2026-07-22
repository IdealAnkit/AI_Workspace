"""Mock-backed tests for the synchronous document processing foundation."""

import io
import unittest
import uuid
from datetime import UTC, datetime
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient

from app.api.deps import get_current_user
from app.core.app_factory import create_app
from app.models.user import User
from app.modules.documents.events.base import DocumentUploadedEvent
from app.modules.documents.models.document import Document, DocumentStatus
from app.modules.document_processing.dependencies import get_processing_service as processing_service_dependency
from app.modules.document_processing.events.consumer import DocumentProcessingEventConsumer
from app.modules.document_processing.parsers.docx_parser import DocxParser
from app.modules.document_processing.parsers.markdown_parser import MarkdownParser
from app.modules.document_processing.parsers.pdf_parser import PdfParser
from app.modules.document_processing.parsers.registry import ParserRegistry
from app.modules.document_processing.parsers.txt_parser import TextParser
from app.modules.document_processing.pipeline.pipeline import DocumentProcessingPipeline
from app.modules.document_processing.processors.normalizer import TextNormalizer
from app.modules.document_processing.extractors.metadata import MetadataExtractor
from app.modules.document_processing.repositories.processing_repository import ProcessingRepository
from app.modules.document_processing.schemas.parsed import ParsedDocument, ParsedPage
from app.modules.document_processing.services.processing_service import ProcessingService


def make_document(*, status: DocumentStatus = DocumentStatus.STORED) -> Document:
    owner_id = uuid.uuid4()
    return Document(
        id=uuid.uuid4(),
        owner_id=owner_id,
        filename="notes.txt",
        original_filename="notes.txt",
        mime_type="text/plain",
        extension=".txt",
        size_bytes=12,
        checksum_sha256="a" * 64,
        storage_key="documents/test/aa/content.txt",
        storage_provider="memory",
        status=status.value,
        uploaded_by=owner_id,
        updated_by=owner_id,
        version=1,
        is_latest_version=True,
        is_deleted=False,
        download_count=0,
        metadata_={},
    )


class ParserAndPipelineTests(unittest.IsolatedAsyncioTestCase):
    """Verify parser resolution and normalization without external services."""

    async def test_registry_selects_supported_parsers_and_pipeline_normalizes_text(self) -> None:
        registry = ParserRegistry([PdfParser(), TextParser(), MarkdownParser(), DocxParser()])
        self.assertEqual(registry.resolve(make_document()).name, "txt")
        self.assertEqual(registry.resolve(make_document_for("readme.md", "text/markdown", ".md")).name, "markdown")
        self.assertEqual(registry.resolve(make_document_for("a.pdf", "application/pdf", ".pdf")).name, "pdf")
        self.assertEqual(
            registry.resolve(
                make_document_for(
                    "a.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ".docx",
                )
            ).name,
            "docx",
        )

        document = make_document()
        pipeline = DocumentProcessingPipeline(registry, TextNormalizer(), MetadataExtractor())
        parsed = await pipeline.process(document, b"Hello\r\n\r\n\r\nWorld\x00")

        self.assertEqual(parsed.text, "Hello\n\nWorld")
        self.assertEqual(parsed.metadata["word_count"], 2)
        self.assertEqual(parsed.page_count, 1)

    async def test_markdown_docx_and_pdf_extractors_return_common_shape(self) -> None:
        markdown_document = make_document_for("guide.md", "text/markdown", ".md")
        markdown = MarkdownParser().extract(markdown_document.id, b"# Guide\n\nHello", markdown_document)
        self.assertEqual(markdown.title, "Guide")
        self.assertEqual(markdown.pages[0].text, "# Guide\n\nHello")

        from docx import Document as DocxDocument

        source = DocxDocument()
        source.core_properties.title = "DOCX Guide"
        source.add_paragraph("First paragraph")
        buffer = io.BytesIO()
        source.save(buffer)
        docx_document = make_document_for(
            "guide.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".docx",
        )
        docx = DocxParser().extract(docx_document.id, buffer.getvalue(), docx_document)
        self.assertEqual(docx.title, "DOCX Guide")
        self.assertIn("First paragraph", docx.text)

        pdf_document = make_document_for("guide.pdf", "application/pdf", ".pdf")
        page = MagicMock()
        page.extract_text.return_value = "PDF text"
        pdf = MagicMock()
        pdf.metadata = {"Title": "PDF Guide", "Producer": "test"}
        pdf.pages = [page]
        context = MagicMock()
        context.__enter__.return_value = pdf
        with patch("app.modules.document_processing.parsers.pdf_parser.pdfplumber.open", return_value=context):
            parsed_pdf = PdfParser().extract(pdf_document.id, b"%PDF-mock", pdf_document)
        self.assertEqual(parsed_pdf.title, "PDF Guide")
        self.assertEqual(parsed_pdf.pages[0].text, "PDF text")


def make_document_for(filename: str, mime_type: str, extension: str) -> Document:
    document = make_document()
    document.filename = filename
    document.mime_type = mime_type
    document.extension = extension
    return document


class InMemoryDocumentRepository:
    def __init__(self, document: Document) -> None:
        self.document = document

    async def get_by_id(self, document_id: uuid.UUID) -> Document | None:
        return self.document if self.document.id == document_id else None

    async def get_by_id_for_owner(self, document_id: uuid.UUID, owner_id: uuid.UUID) -> Document | None:
        if self.document.id == document_id and self.document.owner_id == owner_id:
            return self.document
        return None

    async def transition_status(self, document: Document, status: DocumentStatus, updated_by: uuid.UUID | None) -> None:
        document.status = status.value
        document.updated_by = updated_by


class InMemoryProcessingRepository:
    def __init__(self) -> None:
        self.job = SimpleNamespace(id=uuid.uuid4(), status="queued")
        self.logs: list[str] = []
        self.errors: list[str] = []
        self.result = None

    async def create_job(self, document_id: uuid.UUID):
        self.job.document_id = document_id
        return self.job

    async def add_log(self, job_id: uuid.UUID, level: str, message: str):
        self.logs.append(message)

    async def set_job_processing(self, job, started_at: datetime):
        job.status = "processing"
        job.started_at = started_at

    async def upsert_result(self, parsed, job_id: uuid.UUID):
        self.result = parsed
        return parsed

    async def set_job_completed(self, job, finished_at: datetime, duration: float):
        job.status = "completed"
        job.finished_at = finished_at

    async def set_job_failed(self, job, finished_at: datetime, duration: float, message: str):
        job.status = "failed"
        job.error_message = message

    async def add_error(self, job_id: uuid.UUID, error_type: str, message: str):
        self.errors.append(message)


class FixedReader:
    async def read(self, document: Document) -> bytes:
        return b"content"


class FixedPipeline:
    async def process(self, document: Document, content: bytes) -> ParsedDocument:
        return ParsedDocument(
            document_id=document.id,
            page_count=1,
            pages=[ParsedPage(number=1, text="content")],
            text="content",
            metadata={},
        )


class FailingPipeline:
    async def process(self, document: Document, content: bytes) -> ParsedDocument:
        raise ValueError("bad source")


class ProcessingServiceTests(unittest.IsolatedAsyncioTestCase):
    """Ensure events drive the document lifecycle and preserve failures."""

    async def test_upload_event_drives_successful_processing_lifecycle(self) -> None:
        document = make_document()
        repository = InMemoryProcessingRepository()
        service = ProcessingService(InMemoryDocumentRepository(document), repository, FixedReader(), FixedPipeline())  # type: ignore[arg-type]
        event = DocumentUploadedEvent(document.id, document.owner_id, datetime.now(UTC))

        await DocumentProcessingEventConsumer(service).handle(event)

        self.assertEqual(document.status, DocumentStatus.READY.value)
        self.assertEqual(repository.job.status, "completed")
        self.assertIsNotNone(repository.result)
        self.assertIn("Processing completed successfully.", repository.logs)

    async def test_processing_failure_sets_failed_state_and_persists_error(self) -> None:
        document = make_document()
        repository = InMemoryProcessingRepository()
        service = ProcessingService(InMemoryDocumentRepository(document), repository, FixedReader(), FailingPipeline())  # type: ignore[arg-type]

        await service.process_document(document.id)

        self.assertEqual(document.status, DocumentStatus.FAILED.value)
        self.assertEqual(repository.job.status, "failed")
        self.assertEqual(repository.errors, ["bad source"])


class ScalarResult:
    def __init__(self, value):
        self.value = value

    def scalar_one_or_none(self):
        return self.value


class QuerySession:
    def __init__(self) -> None:
        self.statement = None

    async def execute(self, statement):
        self.statement = statement
        return ScalarResult(None)


class ProcessingRepositoryTests(unittest.IsolatedAsyncioTestCase):
    async def test_owner_scoped_job_lookup_excludes_soft_deleted_documents(self) -> None:
        session = QuerySession()
        repository = ProcessingRepository(session)  # type: ignore[arg-type]
        await repository.get_job_for_owner(uuid.uuid4(), uuid.uuid4())
        self.assertIn("documents.is_deleted IS false", str(session.statement))


class ApiProcessingService:
    async def list_jobs(self, owner_id, offset, limit):
        return [self._job()], 1

    async def get_job(self, job_id, owner_id):
        return self._job(job_id)

    async def get_logs(self, job_id, owner_id):
        return [SimpleNamespace(id=uuid.uuid4(), processing_job_id=job_id, level="INFO", message="Done", created_at=datetime.now(UTC))]

    async def get_document_overview(self, document_id, owner_id):
        return SimpleNamespace(job=self._job(document_id=document_id), result=None)

    @staticmethod
    def _job(job_id=None, document_id=None):
        now = datetime.now(UTC)
        return SimpleNamespace(
            id=job_id or uuid.uuid4(), document_id=document_id or uuid.uuid4(), status="completed",
            started_at=now, finished_at=now, duration_seconds=0.1, error_message=None,
            attempts=1, created_at=now, updated_at=now,
        )


class ProcessingApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.app = create_app()
        self.user = User(id=uuid.uuid4())
        self.service = ApiProcessingService()
        self.app.dependency_overrides[get_current_user] = lambda: self.user
        self.app.dependency_overrides[processing_service_dependency] = lambda: self.service

    def tearDown(self) -> None:
        self.app.dependency_overrides.clear()

    def test_processing_endpoints_require_authenticated_service_and_use_envelopes(self) -> None:
        with TestClient(self.app) as client:
            jobs = client.get("/api/v1/processing/jobs")
            job_id = jobs.json()["data"][0]["id"]
            detail = client.get(f"/api/v1/processing/jobs/{job_id}")
            logs = client.get(f"/api/v1/processing/jobs/{job_id}/logs")
            overview = client.get(f"/api/v1/documents/{uuid.uuid4()}/processing")

        self.assertEqual(jobs.status_code, 200)
        self.assertTrue(jobs.json()["success"])
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(logs.json()["data"][0]["level"], "INFO")
        self.assertEqual(overview.status_code, 200)
