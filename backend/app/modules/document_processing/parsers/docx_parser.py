"""DOCX parser powered by python-docx."""

import io
from pathlib import Path
from uuid import UUID

from docx import Document as DocxDocument

from app.modules.documents.models.document import Document
from app.modules.document_processing.parsers.base import BaseParser
from app.modules.document_processing.schemas.parsed import ParsedDocument, ParsedPage


class DocxParser(BaseParser):
    """Extract paragraph text and core document properties from DOCX files."""

    name = "docx"

    def supports(self, document: Document) -> bool:
        return (
            document.mime_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            and document.extension == ".docx"
        )

    def extract(self, document_id: UUID, content: bytes, document: Document) -> ParsedDocument:
        docx_document = DocxDocument(io.BytesIO(content))
        text = "\n\n".join(
            paragraph.text.strip() for paragraph in docx_document.paragraphs if paragraph.text.strip()
        )
        properties = docx_document.core_properties
        title = properties.title or Path(document.filename).stem
        return ParsedDocument(
            document_id=document_id,
            title=title,
            page_count=None,
            pages=[ParsedPage(number=1, text=text)],
            text=text,
            metadata=self.metadata(content, document, properties),
            warnings=[],
        )

    def metadata(self, content: bytes, document: Document, properties=None) -> dict:
        if properties is None:
            return {"format": "docx", "source_size_bytes": len(content)}
        return {
            "format": "docx",
            "source_size_bytes": len(content),
            "author": properties.author,
            "created_at": properties.created.isoformat() if properties.created else None,
            "modified_at": properties.modified.isoformat() if properties.modified else None,
        }
